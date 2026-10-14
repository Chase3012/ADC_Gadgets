import os
import sys
import json
import base64
import time
import shutil
import urllib.request
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tc_pr.append(tc_mar)

# 13 Detailed Page Flowchart Definitions
FLOWCHARTS = [
    {
        "id": "User_Login_Flowchart",
        "title": "User Login and Authentication Flowchart",
        "page_file": "src/user/login.html",
        "portal": "User Portal (Customer Side)",
        "summary": "Handles customer authentication, session token generation, error feedback, and password recovery redirection.",
        "mermaid": """graph TD
    Start(["Start: Customer visits login.html"]) --> Choice{"User Action"}
    Choice -->|Login Attempt| Input["Enter Email and Password"]
    Input --> CheckAuth{"Credentials Valid?"}
    CheckAuth -->|Yes| SaveSession["Generate Session Token and Save State"]
    SaveSession --> ConnHome(("To user-home.html"))
    ConnHome --> EndSuccess(["End: Session Authenticated"])
    CheckAuth -->|No| ErrMsg["Display Error: Invalid Credentials"]
    ErrMsg --> EndFail(["End: Authentication Failed"])
    Choice -->|Forgot Password| EnterResetEmail["Enter Registered Account Email"]
    EnterResetEmail --> SendLink["Send Password Reset Link via Email"]
    SendLink --> EndReset(["End: Recovery Email Sent"])""",
        "connectors": [
            ("((To user-home.html))", "Outgoing Connector", "Navigates to Customer Homepage upon successful credential verification.")
        ],
        "steps": [
            ("Start", "Terminal", "Customer navigates to login.html directly or via redirection."),
            ("User Action", "Decision", "Customer selects either standard login or password recovery."),
            ("Enter Credentials", "Input / Process", "Customer inputs registered email and password."),
            ("Credentials Valid?", "Decision", "Backend verifies bcrypt hash in PostgreSQL users table."),
            ("Generate Session Token", "Process", "Issues secure session token stored in browser localStorage."),
            ("To user-home.html", "Connector", "Transfers execution flow to customer homepage (user-home.html)."),
            ("Forgot Password Flow", "Alternate Path", "Sends secure password reset link to user's registered email."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "User_Home_Flowchart",
        "title": "Customer Homepage and Discovery Flowchart",
        "page_file": "src/user/user-home.html",
        "portal": "User Portal (Customer Side)",
        "summary": "Main dashboard for customers to view active store announcements, explore featured devices, and navigate to catalog, account, or support.",
        "mermaid": """graph TD
    ConnIn(("From login.html")) --> Start(["Start: Customer on user-home.html"])
    Start --> FetchData["Fetch Store Announcements and Featured Gadgets"]
    FetchData --> RenderUI["Display Hero Banner, Announcement Ticker and Carousel"]
    RenderUI --> Choice{"Select Action"}
    Choice -->|Browse Catalog| ConnCatalog(("To devices.html"))
    ConnCatalog --> EndNode(["End"])
    Choice -->|My Account| ConnAccount(("To account.html"))
    ConnAccount --> EndNode
    Choice -->|Support Desk| ConnSupport(("To support.html"))
    ConnSupport --> EndNode
    Choice -->|Featured Device| ClickFeatured["Click Featured iPhone Card"]
    ClickFeatured --> QuickSpecs["View Quick Specs Modal"]
    QuickSpecs --> CloseSpecs["Close Modal"]
    CloseSpecs --> EndNode""",
        "connectors": [
            ("((From login.html))", "Incoming Connector", "Customer arrives after successful login."),
            ("((To devices.html))", "Outgoing Connector", "Navigates to full device catalog."),
            ("((To account.html))", "Outgoing Connector", "Navigates to customer account and loan balance tracker."),
            ("((To support.html))", "Outgoing Connector", "Navigates to store FAQ and live chat widget.")
        ],
        "steps": [
            ("Start", "Terminal", "Loads user-home.html dashboard."),
            ("Fetch Store Data", "Process", "Queries PostgreSQL for active announcements and featured gadgets."),
            ("Select Action", "Decision", "User chooses navigation branch."),
            ("To devices.html", "Connector", "Transfers flow to device browsing."),
            ("To account.html", "Connector", "Transfers flow to personal account ledger."),
            ("To support.html", "Connector", "Transfers flow to live support desk."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "User_Devices_Catalog_Flowchart",
        "title": "Device Catalog and Specification Viewer Flowchart",
        "page_file": "src/user/devices.html",
        "portal": "User Portal (Customer Side)",
        "summary": "Allows customers to filter available vs. upcoming iPhones, search by model name, and inspect authentic AppleDB studio renders and specs.",
        "mermaid": """graph TD
    ConnIn(("From user-home.html")) --> Start(["Start: Customer on devices.html"])
    Start --> LoadDB["Fetch Devices from PostgreSQL Database"]
    LoadDB --> FilterChoice{"User Filter Choice"}
    FilterChoice -->|Available Tab| ShowAvail["Render In-Stock iPhones"]
    FilterChoice -->|Upcoming Tab| ShowUpcom["Render Coming Soon iPhones"]
    FilterChoice -->|Search Input| LiveFilter["Filter Grid in Real-Time by Name"]
    ShowAvail --> SelectCard["Click Device Card"]
    ShowUpcom --> SelectCard
    LiveFilter --> SelectCard
    SelectCard --> OpenModal["Open Product Modal with AppleDB Studio Render and Specs"]
    OpenModal --> ModalAction{"Action in Modal"}
    ModalAction -->|Inquire via Chat| ConnChat(("To support.html"))
    ConnChat --> EndNode(["End"])
    ModalAction -->|Close Modal| CloseM["Close Modal and Return to Grid"]
    CloseM --> EndNode""",
        "connectors": [
            ("((From user-home.html))", "Incoming Connector", "User enters catalog from homepage navigation."),
            ("((To support.html))", "Outgoing Connector", "Transfers user to live support chat to inquire about device purchase.")
        ],
        "steps": [
            ("Start", "Terminal", "Initializes devices.html catalog page."),
            ("Fetch Devices", "Process", "Queries database for all iPhone records and AppleDB keys."),
            ("Filter Choice", "Decision", "Customer toggles between Available / Upcoming tabs or searches."),
            ("Click Device Card", "Process", "Launches specification modal for chosen model."),
            ("Inspect AppleDB Specs", "Process", "Displays official studio render, chip, display, camera, and SRP."),
            ("To support.html", "Connector", "Directs customer to live chat for store purchase inquiries."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "User_Account_Flowchart",
        "title": "Customer Account and Loan Ledger Flowchart",
        "page_file": "src/user/account.html",
        "portal": "User Portal (Customer Side)",
        "summary": "Customer personal portal to monitor assigned active device loans (provisioned face-to-face by admin), remaining balance, and amortization history.",
        "mermaid": """graph TD
    ConnIn(("From user-home.html")) --> Start(["Start: Customer on account.html"])
    Start --> CheckSession{"Session Valid?"}
    CheckSession -->|No| ConnLogin(("To login.html"))
    ConnLogin --> EndNode(["End"])
    CheckSession -->|Yes| QueryLoan["Query Assigned Loans and Profile from DB"]
    QueryLoan --> HasActiveLoan{"Active Loan Assigned?"}
    HasActiveLoan -->|Yes| RenderLedger["Display Device Model, SRP, Paid Balance, Remaining Balance and Due Date"]
    HasActiveLoan -->|No| RenderNotice["Display: No Active Loans on File"]
    RenderLedger --> UserChoice{"Account Action"}
    RenderNotice --> UserChoice
    UserChoice -->|Update Password| ChangePass["Update Password and Save to DB"]
    ChangePass --> EndNode
    UserChoice -->|Logout| DoLogout["Clear Token and Session State"]
    DoLogout --> ConnLogout(("To login.html"))
    ConnLogout --> EndNode
    UserChoice -->|Return Home| ConnHome(("To user-home.html"))
    ConnHome --> EndNode""",
        "connectors": [
            ("((From user-home.html))", "Incoming Connector", "Arrives from homepage navigation."),
            ("((To login.html))", "Outgoing Connector", "Redirects to login if session expires or user logs out."),
            ("((To user-home.html))", "Outgoing Connector", "Returns customer to store homepage.")
        ],
        "steps": [
            ("Start", "Terminal", "Loads account.html dashboard."),
            ("Check Session", "Decision", "Validates browser authentication token."),
            ("Query Assigned Loans", "Process", "Fetches active loans created by admin during face-to-face agreement."),
            ("Render Ledger", "Process", "Displays remaining balance, monthly amortization, and schedule."),
            ("Logout / Exit", "Connector", "Transfers flow to login.html or user-home.html."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "User_Support_Chat_Flowchart",
        "title": "Customer Support and Live Chat Flowchart",
        "page_file": "src/user/support.html",
        "portal": "User Portal (Customer Side)",
        "summary": "Provides store FAQs (branch locations, payment channels, warranty) and real-time live messaging with ADC Gadgets staff.",
        "mermaid": """graph TD
    ConnIn(("From devices.html or user-home.html")) --> Start(["Start: Customer on support.html"])
    Start --> LoadSupport["Load Support Knowledgebase and Chat Widget"]
    LoadSupport --> SupportChoice{"Select Support Option"}
    SupportChoice -->|Read Store FAQs| ViewFAQs["Browse Payment Modes, Branch Info and Warranty FAQs"]
    ViewFAQs --> EndNode(["End"])
    SupportChoice -->|Live Chat| OpenWidget["Open Live Chat Widget"]
    OpenWidget --> TypeInquiry["Type Inquiries on Color Availability or Payment"]
    TypeInquiry --> SendMsg["Send Message to Admin Support Desk"]
    SendMsg --> WaitStaff["Wait for Staff Real-Time Response"]
    WaitStaff --> ReceiveReply["Receive Notification and Staff Guidance"]
    ReceiveReply --> CloseChat["Close Chat or Continue Conversation"]
    CloseChat --> EndNode""",
        "connectors": [
            ("((From devices.html or user-home.html))", "Incoming Connector", "Arrives from catalog inquiry or navbar.")
        ],
        "steps": [
            ("Start", "Terminal", "Customer opens support.html or launches floating chat widget."),
            ("Support Choice", "Decision", "Chooses self-service FAQs or live messaging."),
            ("Send Message", "Process", "Transmits inquiry to Express server / database."),
            ("Receive Reply", "Process", "Staff reply delivered via real-time WebSocket / polling."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Login_Flowchart",
        "title": "Administrator Authentication Flowchart",
        "page_file": "src/admin/login.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Authenticates administrative personnel and establishes authorized session before granting dashboard access.",
        "mermaid": """graph TD
    Start(["Start: Administrator visits admin/login.html"]) --> InputCreds["Enter Admin Email and Password"]
    InputCreds --> VerifyAdmin{"Role and Credentials Valid?"}
    VerifyAdmin -->|Yes| IssueAdminToken["Set Admin Session Token"]
    IssueAdminToken --> ConnDash(("To admin-dashboard.html"))
    ConnDash --> EndSuccess(["End: Admin Authenticated"])
    VerifyAdmin -->|No| ShowDenied["Display Error: Access Denied or Invalid Credentials"]
    ShowDenied --> EndFail(["End: Access Denied"])""",
        "connectors": [
            ("((To admin-dashboard.html))", "Outgoing Connector", "Navigates to Executive Dashboard upon successful authentication.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens admin/login.html."),
            ("Enter Admin Credentials", "Input", "Enters administrative email and password."),
            ("Verify Role and Hash", "Decision", "Verifies credentials against admin accounts in database."),
            ("Issue Token", "Process", "Stores admin session token."),
            ("To admin-dashboard.html", "Connector", "Transfers flow to Admin Command Center."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Dashboard_Flowchart",
        "title": "Admin Executive Dashboard Flowchart",
        "page_file": "src/admin/dashboard.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Executive command center summarizing operational KPIs and providing routing to all management modules.",
        "mermaid": """graph TD
    ConnIn(("From admin/login.html")) --> Start(["Start: Admin on admin/dashboard.html"])
    Start --> FetchKPIs["Fetch Operational Metrics: Total Revenue, Active Loans, Low Stock"]
    FetchKPIs --> RenderDashboard["Display KPI Metric Cards and Activity Stream"]
    RenderDashboard --> MenuChoice{"Select Management Module"}
    MenuChoice -->|User and Loan Management| ConnUsers(("To manage-users.html"))
    ConnUsers --> EndNode(["End"])
    MenuChoice -->|Inventory Control| ConnInv(("To inventory.html"))
    ConnInv --> EndNode
    MenuChoice -->|Loan Ledger and Payments| ConnLedger(("To loan-ledger.html"))
    ConnLedger --> EndNode
    MenuChoice -->|Customer Chat Desk| ConnComm(("To communication.html"))
    ConnComm --> EndNode
    MenuChoice -->|Analytics and Reports| ConnReports(("To reports.html"))
    ConnReports --> EndNode
    MenuChoice -->|Staff Administration| ConnStaff(("To manage-admins.html"))
    ConnStaff --> EndNode
    MenuChoice -->|Logout| ClearSession["Clear Admin Session"]
    ClearSession --> ConnLogout(("To admin/login.html"))
    ConnLogout --> EndNode""",
        "connectors": [
            ("((From admin/login.html))", "Incoming Connector", "Arrives after successful admin login."),
            ("((To manage-users.html))", "Outgoing Connector", "Transfers to Customer and Loan provisioning."),
            ("((To inventory.html))", "Outgoing Connector", "Transfers to Device Inventory management."),
            ("((To loan-ledger.html))", "Outgoing Connector", "Transfers to Payment Ledger and Collections."),
            ("((To communication.html))", "Outgoing Connector", "Transfers to Live Customer Support desk."),
            ("((To reports.html))", "Outgoing Connector", "Transfers to Financial and Business Reports."),
            ("((To manage-admins.html))", "Outgoing Connector", "Transfers to Admin Staff management.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens dashboard."),
            ("Fetch KPIs", "Process", "Aggregates revenue, active loans, and inventory counts."),
            ("Select Module", "Decision", "Admin selects administrative feature to manage."),
            ("Module Connectors", "Connector", "Transfers flow to specific administrative page."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Manage_Users_and_Loan_Flowchart",
        "title": "Customer Account and Direct Loan Assignment Flowchart",
        "page_file": "src/admin/manage-users.html & user-loan.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Core administrative workflow: provisions customer accounts and registers verified face-to-face loan agreements onto the system.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Admin on manage-users.html"])
    Start --> WalkInNotice["Customer Completes Face-to-Face Agreement at Branch"]
    WalkInNotice --> AdminTask{"Select Action"}
    AdminTask -->|View Existing Users| SearchUsers["Search Customer by Name or Email"]
    SearchUsers --> InspectProfile["View Customer Contact Details and Active Loans"]
    InspectProfile --> EndNode(["End"])
    AdminTask -->|Create User and Assign Loan| OpenCreateModal["Click Create User and Assign Loan"]
    OpenCreateModal --> InputProfile["Enter Name, Email, Contact and Address"]
    InputProfile --> SelectDevice["Select In-Stock iPhone Model from Inventory"]
    SelectDevice --> SetTerms["Input SRP, Downpayment and Loan Term: 6, 12, or 24 Mos"]
    SetTerms --> CalcAmort["System Calculates Monthly Amortization and Payment Schedule"]
    CalcAmort --> SubmitRecord["Submit to PostgreSQL Database"]
    SubmitRecord --> WriteDB["Write User Account + Insert Loan Record + Deduct 1 Stock"]
    WriteDB --> IssueCreds["Display Generated Credentials to Hand to Customer"]
    IssueCreds --> EndNode""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens customer provisioning module."),
            ("Face-to-Face Agreement", "External Trigger", "Customer physically completes agreement at store."),
            ("Input Profile", "Input", "Admin enters customer demographic details."),
            ("Select Device and Terms", "Process", "Admin picks in-stock device and defines term length."),
            ("System Calculation", "Process", "Automated amortization calculation based on term."),
            ("Write Database", "Database Write", "Creates user record, inserts loan contract, decrements stock."),
            ("Issue Credentials", "Output", "Admin gives login credentials to customer for account tracking."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Inventory_Flowchart",
        "title": "Inventory Control and AppleDB Integration Flowchart",
        "page_file": "src/admin/inventory.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Manages device catalog, utilizing AppleDB search to auto-populate hardware specs and official studio renders, and manages stock levels.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Admin on inventory.html"])
    Start --> LoadInventory["Fetch Device Inventory from Database"]
    LoadInventory --> InvChoice{"Inventory Action"}
    InvChoice -->|Add Device via AppleDB| OpenAddModal["Click Add Device"]
    OpenAddModal --> SearchADB["Type Model Name in AppleDB Search"]
    SearchADB --> AutoFill["Auto-Populate SoC Chip, Colors and Official Studio Render"]
    AutoFill --> InputStockPrice["Enter SRP, Monthly Factor and Stock Count"]
    AutoFill --> SetStatus["Set Status: Available or Upcoming"]
    SetStatus --> SaveDevice["Insert New Device into Database"]
    SaveDevice --> EndNode(["End"])
    InvChoice -->|Edit Device| SelectRow["Select Device Row in Table"]
    SelectRow --> EditModal["Adjust Stock Count, Status or Pricing"]
    EditModal --> UpdateDB["Update Record in Database"]
    UpdateDB --> EndNode
    InvChoice -->|Delete Device| DeleteConfirm["Confirm Deletion of Discontinued Model"]
    DeleteConfirm --> RemoveDB["Delete Record from Database"]
    RemoveDB --> EndNode""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens inventory.html."),
            ("Inventory Action", "Decision", "Chooses Add Device, Edit Stock, or Delete Model."),
            ("AppleDB Auto-Fill", "External Integration", "Queries AppleDB API for verified specs and studio renders."),
            ("Set Stock and Status", "Input", "Sets inventory units and marks Available or Upcoming."),
            ("Database Write", "Database Write", "Commits device changes to PostgreSQL table."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Loan_Ledger_Flowchart",
        "title": "Loan Ledger and Collections Flowchart",
        "page_file": "src/admin/loan-ledger.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Records customer monthly installment payments (Cash, GCash, Bank Transfer) and automatically computes remaining balance and due dates.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Admin on loan-ledger.html"])
    Start --> PayNotice["Customer Submits Monthly Installment Payment"]
    PayNotice --> SearchAccount["Search Customer Name or Loan Account Number"]
    SearchAccount --> ViewBalance["Inspect Outstanding Balance and Current Due Date"]
    ViewBalance --> ClickRecord["Click Record Payment"]
    ClickRecord --> EnterPayment["Enter Amount Paid, Payment Date and Channel: Cash/GCash/Bank"]
    EnterPayment --> SubmitPay["Submit Transaction"]
    SubmitPay --> UpdateDB["Insert Payment Log + Deduct Remaining Balance + Advance Due Date"]
    UpdateDB --> IssueReceipt["Generate Official Payment Receipt for Customer"]
    IssueReceipt --> EndNode(["End"])""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens loan-ledger.html."),
            ("Payment Notice", "External Trigger", "Customer pays installment at store or sends transfer slip."),
            ("Search Account", "Input", "Admin locates active loan record in ledger."),
            ("Record Payment", "Input", "Enters amount paid, timestamp, and payment method."),
            ("Update Database", "Database Write", "Deducts payment from balance and rolls forward due date."),
            ("Issue Receipt", "Output", "Generates official receipt for the borrower."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Communication_Flowchart",
        "title": "Live Support Desk and Messaging Flowchart",
        "page_file": "src/admin/communication.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Unified inbox for staff to review incoming customer inquiries and reply in real time.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Admin on communication.html"])
    Start --> FetchThreads["Load Customer Conversation Threads from DB"]
    FetchThreads --> CheckUnread{"Unread Inquiries?"}
    CheckUnread -->|Yes| HighlightBadge["Display Red Notification Badge on Thread"]
    CheckUnread -->|No| IdleInbox["Display Active Inbox State"]
    HighlightBadge --> OpenThread["Select Customer Conversation Thread"]
    IdleInbox --> OpenThread
    OpenThread --> ReadMsg["Read Customer Question: Stock, Price, or Payment Proof"]
    ReadMsg --> TypeReply["Type Staff Guidance Response"]
    TypeReply --> SendReply["Click Send Message"]
    SendReply --> PushDB["Write Message to DB and Push to Customer Widget in Real-Time"]
    PushDB --> EndNode(["End"])""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens communication.html."),
            ("Check Unread", "Decision", "Identifies threads with pending customer messages."),
            ("Read Message", "Process", "Admin reviews customer inquiry."),
            ("Send Reply", "Process", "Sends advice, payment instructions, or stock availability."),
            ("Push to DB", "Database Write", "Real-time sync to customer's live chat session."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Reports_Flowchart",
        "title": "Business Reports and Analytics Flowchart",
        "page_file": "src/admin/reports.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Compiles financial analytics, collection efficiency, loan disbursement volume, and inventory velocity for administrative reporting.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Admin on reports.html"])
    Start --> LoadAnalytics["Load Financial and Sales Analytics Dashboard"]
    LoadAnalytics --> SelectFilters["Select Date Range and Report Category"]
    SelectFilters --> QueryDB["Aggregate PostgreSQL Sales and Payment Data"]
    QueryDB --> RenderCharts["Generate Visual Revenue, Collections and Inventory Charts"]
    RenderCharts --> ExportChoice{"Export Report?"}
    ExportChoice -->|Export CSV| GenCSV["Download Spreadsheet Data in CSV Format"]
    GenCSV --> EndNode(["End"])
    ExportChoice -->|Export PDF| GenPDF["Generate Print-Ready Audit PDF Document"]
    GenPDF --> EndNode
    ExportChoice -->|View Only| EndNode""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Admin opens reports.html."),
            ("Select Filters", "Input", "Chooses date range (daily, weekly, monthly) and category."),
            ("Aggregate DB Data", "Process", "Queries database for financial summaries."),
            ("Render Charts", "Output", "Displays visual metrics and summary tables."),
            ("Export Choice", "Decision", "Optionally downloads CSV or PDF for bookkeeping."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    },
    {
        "id": "Admin_Manage_Admins_Flowchart",
        "title": "Staff and Admin Access Control Flowchart",
        "page_file": "src/admin/manage-admins.html",
        "portal": "Admin Portal (Management Side)",
        "summary": "Security administration module to create, manage, and revoke internal staff credentials.",
        "mermaid": """graph TD
    ConnIn(("From admin/dashboard.html")) --> Start(["Start: Super Admin on manage-admins.html"])
    Start --> FetchStaff["Fetch List of Active Staff and Admin Accounts"]
    FetchStaff --> RenderTable["Display Staff Name, Email, Role and Account Status"]
    RenderTable --> StaffAction{"Select Staff Action"}
    StaffAction -->|Add New Staff| OpenStaffModal["Click Add Staff Member"]
    OpenStaffModal --> InputStaffData["Enter Name, Work Email, Role and Initial Password"]
    InputStaffData --> SaveStaff["Write New Staff Credentials to Database"]
    SaveStaff --> EndNode(["End"])
    StaffAction -->|Revoke Access| SelectStaffRow["Locate Staff Account Row"]
    SelectStaffRow --> ConfirmRevoke["Confirm Deactivation"]
    ConfirmRevoke --> UpdateStatus["Set Account Status to Inactive / Revoked in DB"]
    UpdateStatus --> EndNode""",
        "connectors": [
            ("((From admin/dashboard.html))", "Incoming Connector", "Arrives from dashboard menu.")
        ],
        "steps": [
            ("Start", "Terminal", "Super Admin opens manage-admins.html."),
            ("Staff Action", "Decision", "Chooses to provision new staff or revoke existing access."),
            ("Input Staff Data", "Input", "Enters work email, security role, and password."),
            ("Save to Database", "Database Write", "Creates staff record in admin_users table."),
            ("Revoke Access", "Process", "Deactivates staff credentials to prevent system access."),
            ("End", "Terminal", "Terminates current page flow.")
        ]
    }
]

def download_mermaid_image(mermaid_code, output_png_path):
    payload = {
        "code": mermaid_code,
        "mermaid": {"theme": "default"}
    }
    json_bytes = json.dumps(payload).encode('utf-8')
    b64 = base64.b64encode(json_bytes).decode('ascii')
    url = f"https://mermaid.ink/img/{b64}"
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    
    for attempt in range(4):
        try:
            with urllib.request.urlopen(req, timeout=20) as response:
                data = response.read()
                with open(output_png_path, 'wb') as f:
                    f.write(data)
            return True
        except Exception as e:
            time.sleep(2)
            if attempt == 3:
                print(f"Error downloading image for {output_png_path}: {e}")
                return False
    return False

def build_individual_docx(flow_item, png_path, output_docx_path):
    doc = docx.Document()

    # 0.8 inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(2)
    r_title = title_p.add_run(flow_item["title"])
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(30, 27, 75) # Deep Navy

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(10)
    r_sub = sub_p.add_run(f"System: ADC Gadgets  |  Portal: {flow_item['portal']}  |  Page File: {flow_item['page_file']}")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(236, 72, 153) # ADC Pink

    # Scope / Functional Description Callout
    desc_tbl = doc.add_table(rows=1, cols=1)
    desc_cell = desc_tbl.rows[0].cells[0]
    set_cell_background(desc_cell, "F8FAFC")
    set_cell_margins(desc_cell, top=80, bottom=80, left=120, right=120)
    p_desc = desc_cell.paragraphs[0]
    p_desc.paragraph_format.space_after = Pt(2)
    r_dlabel = p_desc.add_run("FUNCTIONAL SUMMARY & SCOPE:\n")
    r_dlabel.font.bold = True
    r_dlabel.font.size = Pt(9.5)
    r_dlabel.font.color.rgb = RGBColor(71, 85, 105)
    r_dtext = p_desc.add_run(flow_item["summary"] + "\n" + 
        "Note: In accordance with system directives, there are no online loan requests. " +
        "Admin provisions user accounts directly from face-to-face agreements.")
    r_dtext.font.size = Pt(9)
    r_dtext.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Visual Flowchart Section
    sec_head = doc.add_paragraph()
    sec_head.paragraph_format.space_before = Pt(8)
    sec_head.paragraph_format.space_after = Pt(4)
    r_sh = sec_head.add_run("1. Flowchart Diagram (Paper-Ready Visual)")
    r_sh.font.name = "Segoe UI"
    r_sh.font.bold = True
    r_sh.font.size = Pt(12)
    r_sh.font.color.rgb = RGBColor(30, 27, 75)

    # Embed PNG Diagram
    if os.path.exists(png_path) and os.path.getsize(png_path) > 0:
        img_p = doc.add_paragraph()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.paragraph_format.space_before = Pt(4)
        img_p.paragraph_format.space_after = Pt(8)
        # Sizing image to fit on standard 8.5x11 page with margins
        img_p.add_run().add_picture(png_path, width=Inches(5.6))
    else:
        warn_p = doc.add_paragraph()
        warn_p.add_run("[Image diagram will be rendered and attached]").font.color.rgb = RGBColor(100, 116, 139)

    # Connected Pages / Connectors Table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    con_head = doc.add_paragraph()
    con_head.paragraph_format.space_before = Pt(8)
    con_head.paragraph_format.space_after = Pt(4)
    r_con = con_head.add_run("2. Page Connectors (Navigation Relationships)")
    r_con.font.name = "Segoe UI"
    r_con.font.bold = True
    r_con.font.size = Pt(12)
    r_con.font.color.rgb = RGBColor(30, 27, 75)

    con_tbl = doc.add_table(rows=len(flow_item["connectors"]) + 1, cols=3)
    con_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = con_tbl.rows[0].cells
    hdr[0].width = Inches(2.0)
    hdr[1].width = Inches(1.8)
    hdr[2].width = Inches(3.0)
    for i, title in enumerate(["Connector Symbol", "Connector Type", "Target Page & Explanation"]):
        set_cell_background(hdr[i], "1E1B4B")
        set_cell_margins(hdr[i], top=60, bottom=60, left=80, right=80)
        p = hdr[i].paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (sym, ctype, cdesc) in enumerate(flow_item["connectors"]):
        row_cells = con_tbl.rows[idx + 1].cells
        row_cells[0].width = Inches(2.0)
        row_cells[1].width = Inches(1.8)
        row_cells[2].width = Inches(3.0)
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([sym, ctype, cdesc]):
            set_cell_background(row_cells[i], bg)
            set_cell_margins(row_cells[i], top=50, bottom=50, left=80, right=80)
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(190, 18, 60)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    # Step-by-Step Logic Table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    step_head = doc.add_paragraph()
    step_head.paragraph_format.space_before = Pt(8)
    step_head.paragraph_format.space_after = Pt(4)
    r_step = step_head.add_run("3. Step-by-Step Process & Logic Table")
    r_step.font.name = "Segoe UI"
    r_step.font.bold = True
    r_step.font.size = Pt(12)
    r_step.font.color.rgb = RGBColor(30, 27, 75)

    step_tbl = doc.add_table(rows=len(flow_item["steps"]) + 1, cols=3)
    step_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    shdr = step_tbl.rows[0].cells
    shdr[0].width = Inches(2.2)
    shdr[1].width = Inches(1.5)
    shdr[2].width = Inches(3.1)
    for i, title in enumerate(["Flowchart Step / Node", "Flowchart Shape", "Logic & Action Description"]):
        set_cell_background(shdr[i], "1E1B4B")
        set_cell_margins(shdr[i], top=60, bottom=60, left=80, right=80)
        p = shdr[i].paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for idx, (node_name, shape_name, node_desc) in enumerate(flow_item["steps"]):
        row_cells = step_tbl.rows[idx + 1].cells
        row_cells[0].width = Inches(2.2)
        row_cells[1].width = Inches(1.5)
        row_cells[2].width = Inches(3.1)
        bg = "F8FAFC" if idx % 2 == 1 else "FFFFFF"
        for i, val in enumerate([node_name, shape_name, node_desc]):
            set_cell_background(row_cells[i], bg)
            set_cell_margins(row_cells[i], top=50, bottom=50, left=80, right=80)
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(8.5)
            if i == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(15, 23, 42)
            elif i == 1:
                r.font.color.rgb = RGBColor(71, 85, 105)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

    # Save Word document
    doc.save(output_docx_path)
    print(f"Generated Word Document: {output_docx_path}")

def main():
    download_dir = r"C:\Users\Chase\Downloads"
    download_subfolder = os.path.join(download_dir, "ADC_Gadgets_Page_Flowcharts")
    workspace_docs = r"c:\Users\Chase\Documents\my-supabase-project\docs\flowcharts"

    os.makedirs(download_dir, exist_ok=True)
    os.makedirs(download_subfolder, exist_ok=True)
    os.makedirs(workspace_docs, exist_ok=True)

    print(f"Starting generation of {len(FLOWCHARTS)} paper-ready flowcharts...")

    for flow in FLOWCHARTS:
        file_base = flow["id"]
        png_temp_path = os.path.join(download_subfolder, f"{file_base}.png")
        
        # 1. Download Diagram PNG from Mermaid renderer if not already present or if small
        needs_download = not os.path.exists(png_temp_path) or os.path.getsize(png_temp_path) < 1000
        if needs_download:
            print(f"\n[+] Downloading diagram: {file_base}")
            success = download_mermaid_image(flow["mermaid"], png_temp_path)
            if not success:
                print(f"Failed to generate diagram for {file_base}")
            time.sleep(1) # respectful pacing for renderer API
        else:
            print(f"\n[+] Using existing diagram: {file_base} ({os.path.getsize(png_temp_path)} bytes)")

        # Also copy PNG to direct Downloads and docs
        if os.path.exists(png_temp_path) and os.path.getsize(png_temp_path) > 1000:
            for dest in [os.path.join(download_dir, f"{file_base}.png"), os.path.join(workspace_docs, f"{file_base}.png")]:
                try:
                    shutil.copyfile(png_temp_path, dest)
                except Exception as e:
                    print(f"Warning copying PNG to {dest}: {e}")

        # 2. Build Word Document (.docx)
        docx_main = os.path.join(download_dir, f"{file_base}.docx")
        docx_sub = os.path.join(download_subfolder, f"{file_base}.docx")
        docx_docs = os.path.join(workspace_docs, f"{file_base}.docx")

        build_individual_docx(flow, png_temp_path, docx_sub)
        
        # Copy to main downloads and docs folder
        try:
            shutil.copyfile(docx_sub, docx_main)
            shutil.copyfile(docx_sub, docx_docs)
        except Exception as e:
            print(f"Warning copying docx: {e}")

    print("\nAll 13 paper-ready flowcharts generated successfully!")

if __name__ == "__main__":
    main()
