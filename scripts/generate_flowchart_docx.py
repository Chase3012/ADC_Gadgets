import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tc_pr.append(tc_mar)

def create_flowchart_document(output_paths):
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # ━━━ COVER / HEADER ━━━
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("ADC GADGETS")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(28)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(30, 27, 75) # Deep Indigo

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("System Flowchart & Process Architecture Documentation (User & Admin Portals)")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(236, 72, 153) # ADC Pink

    # Metadata table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("System Name:", "ADC Gadgets Management & Customer Portal System"),
        ("Architecture:", "HTML5, Vanilla CSS, JavaScript, Express.js, PostgreSQL (Neon)"),
        ("Prepared For:", "Course Project & Academic System Architecture Review"),
        ("Date / Version:", "October 2026 / Version 2.0 (Official AppleDB Edition)")
    ]
    for i, (label, val) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(1.8)
        c1.width = Inches(5.0)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = RGBColor(71, 85, 105)

        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(val)
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ━━━ SECTION 1: IMPORTANT SCOPE DIRECTIVE ━━━
    h1 = doc.add_heading(level=1)
    r_h1 = h1.add_run("1. Executive Summary & Crucial System Scope")
    r_h1.font.color.rgb = RGBColor(30, 27, 75)

    # Callout Note Box
    note_table = doc.add_table(rows=1, cols=1)
    note_cell = note_table.rows[0].cells[0]
    set_cell_background(note_cell, "FFF1F2") # Light Rose
    set_cell_margins(note_cell, top=140, bottom=140, left=180, right=180)
    note_p = note_cell.paragraphs[0]
    r_note_title = note_p.add_run("IMPORTANT ARCHITECTURAL RULE (NO ONLINE LOAN REQUESTS):\n")
    r_note_title.font.bold = True
    r_note_title.font.size = Pt(10.5)
    r_note_title.font.color.rgb = RGBColor(190, 18, 60) # Deep Crimson
    r_note_body = note_p.add_run(
        "There is NO online loan request or credit application form in the ADC Gadgets system. "
        "All loan requests, identity verifications, and financial underwriting agreements are handled "
        "strictly face-to-face at the branch or through external partner platforms.\n\n"
        "• Administrator's Role: To create customer user accounts and directly bind/register their verified "
        "face-to-face device loan contracts onto the system.\n"
        "• Customer's Role: Customers browse verified devices, view official AppleDB studio specifications, "
        "monitor their active loan balances and payment schedules (provisioned by admin), and chat with live support staff."
    )
    r_note_body.font.size = Pt(10)
    r_note_body.font.color.rgb = RGBColor(136, 19, 55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ━━━ SECTION 2: HIGH-LEVEL ARCHITECTURE FLOWCHART ━━━
    h2 = doc.add_heading(level=1)
    r_h2 = h2.add_run("2. High-Level System Architecture Flowchart")
    r_h2.font.color.rgb = RGBColor(30, 27, 75)

    p_desc2 = doc.add_paragraph()
    p_desc2.add_run(
        "This flowchart illustrates the complete relationship between the Face-to-Face Agreement stage, "
        "the Admin Portal, the User Portal, and the central PostgreSQL database:"
    )

    # Mermaid Box
    mermaid_1 = """flowchart TB
    subgraph External["External / Face-to-Face Stage"]
        FTF[Customer Visits Branch / Partner Platform] --> AGREE[Face-to-Face Agreement & Credit Check]
    end

    subgraph AdminPortal["ADMIN PORTAL (Management Side)"]
        ALogin([Admin Login]) --> ADash[Admin Dashboard]
        ADash --> AUsers[User & Loan Management]
        ADash --> AInv[Inventory Management with AppleDB]
        ADash --> ALedger[Loan Ledger & Payments]
        ADash --> AComm[Customer Chat Desk]
        ADash --> ARep[Financial & Sales Reports]
        ADash --> AAdmin[Staff Admin Management]

        AGREE -.->|Walk-in Agreement Details| AUsers
        AUsers -->|Create Account & Assign Loan| DB[(Central PostgreSQL DB)]
        AInv -->|Stock, Pricing, Status| DB
        ALedger -->|Record Monthly Payments| DB
        AComm <-->|Real-time Inquiries| DB
        ARep <---|Fetch Analytics| DB
    end

    subgraph UserPortal["USER PORTAL (Customer Side)"]
        ULogin([Customer Login]) --> UHome[Customer Home Page]
        UHome --> UCatalog[Browse Devices Catalog]
        UHome --> UAccount[My Account & Loan Ledger]
        UHome --> USupport[Support FAQs & Live Chat]

        UCatalog --> UModal[View AppleDB Studio Specs]
        UAccount --> UTrack[View Remaining Balance & Due Dates]
        USupport <--> UChat[Message Staff in Real-time]

        DB -.->|Credentials Provided| ULogin
        DB --->|Fetch Active Loans & Balances| UAccount
        DB --->|Fetch Devices & Announcements| UHome
        DB <--->|Chat Thread Sync| USupport
    end"""

    add_code_block(doc, mermaid_1, "Mermaid Diagram Code (High-Level Architecture)")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ━━━ SECTION 3: USER PORTAL FLOWCHART & PAGES ━━━
    h3 = doc.add_heading(level=1)
    r_h3 = h3.add_run("3. User Portal (Customer Side) Flowchart")
    r_h3.font.color.rgb = RGBColor(30, 27, 75)

    p_desc3 = doc.add_paragraph()
    p_desc3.add_run(
        "The customer portal focuses on product discovery, viewing official hardware specifications, "
        "monitoring active assigned loans and payment due dates, and contacting support:"
    )

    mermaid_2 = """flowchart TD
    StartUser([Customer Starts]) --> VisitSite[Visit ADC Gadgets Website]
    VisitSite --> IsAuth{Logged In?}
    
    IsAuth -- No --> LoginPage[Navigate to login.html]
    LoginPage --> AuthChoice{Action}
    AuthChoice -- Sign In --> EnterCreds[Enter Email & Password]
    AuthChoice -- Forgot Pass --> ForgotPass[Request Password Reset Link]
    ForgotPass --> EnterCreds
    
    EnterCreds --> ValidateAuth{Credentials Valid?}
    ValidateAuth -- No --> AuthErr[Show Error: Invalid Credentials]
    AuthErr --> LoginPage
    ValidateAuth -- Yes --> GenSession[Generate Session & Store Token]
    GenSession --> Home[user-home.html]
    IsAuth -- Yes --> Home

    Home --> UserAction{Select Menu Navigation}

    %% Option A: Browse Catalog
    UserAction -- Browse Catalog --> Catalog[devices.html]
    Catalog --> FilterTab{Select Tab}
    FilterTab -- Available --> ShowAvail[Render In-Stock iPhones]
    FilterTab -- Upcoming --> ShowUpcom[Render Coming Soon iPhones]
    ShowAvail --> SearchDevice[Search by Model Name]
    ShowUpcom --> SearchDevice
    SearchDevice --> ClickCard[Click Device Card]
    ClickCard --> ViewModal[Open Product Specs Modal]
    ViewModal --> InspectSpecs[Inspect AppleDB Specs & Monthly SRP]
    InspectSpecs --> CloseModal[Close Modal or Inquire via Chat]
    CloseModal --> Home

    %% Option B: Account & Loan Ledger
    UserAction -- My Account --> AccountPage[account.html]
    AccountPage --> LoadLoans[Query Assigned Loans from DB]
    LoadLoans --> HasLoan{Has Active Loan?}
    HasLoan -- Yes --> DispLoan[Display Assigned Device, SRP, Paid Amount, Balance, Due Date]
    HasLoan -- No --> DispEmpty[Display: No Active Loans Assigned]
    DispLoan --> ViewSchedule[Inspect Monthly Amortization Schedule]
    ViewSchedule --> AccountPage
    AccountPage --> SecuritySettings[Change Password / Update Profile]
    AccountPage --> LogoutAction[Click Logout]
    LogoutAction --> ClearSession[Clear Local Storage & Session]
    ClearSession --> LoginPage

    %% Option C: Support & Live Chat
    UserAction -- Customer Support --> SupportPage[support.html / Chat Widget]
    SupportPage --> ReadFAQs[Browse Frequently Asked Questions]
    SupportPage --> OpenChat[Open Live Chat Widget]
    OpenChat --> SendMsg[Send Message or Payment Query to Admin]
    SendMsg --> WaitStaff[Wait for Staff Real-Time Response]
    WaitStaff --> ReceiveReply[Receive Notification & Staff Advice]
    ReceiveReply --> SupportPage"""

    add_code_block(doc, mermaid_2, "Mermaid Diagram Code (User Portal Flowchart)")

    # User Pages Table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p_tbl3 = doc.add_paragraph()
    r_tbl3 = p_tbl3.add_run("Customer Portal Page Reference Table:")
    r_tbl3.font.bold = True

    user_pages = [
        ("login.html", "Authentication & Login", "Validates customer login credentials; provides secure session tokens and password recovery."),
        ("user-home.html", "Customer Homepage", "Displays promotional hero banners, real-time store announcement tickers, and featured gadget carousel."),
        ("devices.html", "Device Catalog & Specs", "Interactive browsing with Available / Upcoming filters, live search, and AppleDB hardware specification modal."),
        ("account.html", "Account & Loan Ledger", "Personal customer dashboard displaying profile, assigned active loan records, remaining balance, amortization schedule, and payment history."),
        ("support.html", "Support Desk & Live Chat", "Comprehensive FAQ knowledgebase and real-time live messaging with ADC Gadgets store staff.")
    ]
    add_styled_table(doc, ["Page File", "Module Title", "Functional Description"], user_pages, [1.5, 1.8, 3.5])

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ━━━ SECTION 4: ADMIN PORTAL FLOWCHART & PAGES ━━━
    h4 = doc.add_heading(level=1)
    r_h4 = h4.add_run("4. Admin Portal (Management Side) Flowchart")
    r_h4.font.color.rgb = RGBColor(30, 27, 75)

    p_desc4 = doc.add_paragraph()
    p_desc4.add_run(
        "The Admin Portal flowchart details operational procedures: account provisioning for walk-in customers, "
        "direct loan contract assignment, inventory management with AppleDB synchronization, payment recording, and financial reporting:"
    )

    mermaid_3 = """flowchart TD
    StartAdmin([Administrator Starts]) --> AdminLogin[admin/login.html]
    AdminLogin --> EnterAdminCreds[Enter Admin Email & Password]
    EnterAdminCreds --> CheckAdminAuth{Role & Credentials Valid?}
    
    CheckAdminAuth -- No --> AdminDenied[Display Access Denied]
    AdminDenied --> AdminLogin
    CheckAdminAuth -- Yes --> AdminDash[admin/dashboard.html]

    AdminDash --> ViewKPIs[Inspect KPIs: Total Revenue, Active Loans, Overdue Accounts, Low Stock]
    AdminDash --> AdminMenu{Select Admin Module}

    %% Module 1: User & Loan Provisioning
    AdminMenu -- User & Loan Management --> UserLoanModule[manage-users.html & user-loan.html]
    UserLoanModule --> UserActionChoice{Customer Workflow}
    
    UserActionChoice -- Create User & Assign Loan --> EnterCustomerDetails[Input Customer Name, Email, Phone, Address]
    EnterCustomerDetails --> SelectPhoneModel[Select In-Stock Device from Inventory]
    SelectPhoneModel --> SetLoanTerms[Set SRP, Downpayment & Term: 6, 12, or 24 Mos]
    SetLoanTerms --> CalcAmort[System Computes Monthly Amortization & Due Dates]
    CalcAmort --> SaveCustomerLoan[Commit to DB: User Account + Active Loan Record]
    SaveCustomerLoan --> NotifyClient[Provide Generated Credentials to Customer]
    
    UserActionChoice -- View Customers --> ListClients[Filter & Search Existing Borrowers]
    ListClients --> EditCustomer[Update Contact Info / Status]

    %% Module 2: Inventory Management
    AdminMenu -- Inventory Management --> InvModule[admin/inventory.html]
    InvModule --> InvAction{Inventory Task}
    
    InvAction -- Add Device via AppleDB --> SearchADB[Type iPhone Model in AppleDB Search]
    SearchADB --> AutoPopulate[Auto-Fill Identifier, Soc Chip, Colors, Official Studio Render]
    AutoPopulate --> SetPricingStock[Enter SRP, Monthly Factor, Initial Stock, Status]
    SetPricingStock --> SaveDevice[Insert New Device into DB]
    
    InvAction -- Update Stock / Price --> SelectDeviceRow[Locate Device Row in Inventory Table]
    SelectDeviceRow --> OpenEditModal[Edit Stock Quantity, Status or Price]
    OpenEditModal --> SaveEdit[Update Database]
    
    InvAction -- Remove Discontinued --> DeletePrompt[Confirm Delete Device]
    DeletePrompt --> RemoveDevice[Delete / Archive Device]

    %% Module 3: Loan Ledger & Payments
    AdminMenu -- Loan Ledger & Collections --> LedgerModule[admin/loan-ledger.html]
    LedgerModule --> FindCustomerAccount[Search Customer by Name, Email, or Loan ID]
    FindCustomerAccount --> OpenPaymentModal[Click 'Record Payment']
    OpenPaymentModal --> InputPayment[Enter Amount Paid, Payment Date, Method: Cash/GCash/Bank]
    InputPayment --> ProcessLedger[System Deducts from Remaining Balance & Updates Due Date]
    ProcessLedger --> GenReceipt[Generate Official Payment Ledger Entry]

    %% Module 4: Live Customer Support
    AdminMenu -- Customer Communication --> CommModule[admin/communication.html]
    CommModule --> CheckInbox[View Active Chat Threads]
    CheckInbox --> SelectConversation[Select Customer Conversation]
    SelectConversation --> ReadUserInquiry[Read Inquiry: Product Stock, Payment Confirmation]
    ReadUserInquiry --> SendStaffReply[Type & Send Staff Reply]

    %% Module 5: Financial Reports
    AdminMenu -- Business Reports --> ReportsModule[admin/reports.html]
    ReportsModule --> SelectReportRange[Select Date Range & Report Category]
    SelectReportRange --> RenderCharts[Generate Revenue, Loan Disbursements, and Inventory Turnover]
    RenderCharts --> ExportReport[Export Report to CSV or PDF]

    %% Module 6: Admin Staff Management
    AdminMenu -- Admin Staff Control --> AdminMgtModule[admin/manage-admins.html]
    AdminMgtModule --> ManageStaff[Create New Staff Account or Revoke Access]"""

    add_code_block(doc, mermaid_3, "Mermaid Diagram Code (Admin Portal Flowchart)")

    # Admin Pages Table
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    p_tbl4 = doc.add_paragraph()
    r_tbl4 = p_tbl4.add_run("Admin Portal Page Reference Table:")
    r_tbl4.font.bold = True

    admin_pages = [
        ("admin/login.html", "Admin Authentication", "Authenticates staff and super-admin credentials; verifies security authorization level."),
        ("admin/dashboard.html", "Executive Dashboard", "Consolidated command center displaying total collections, active loan contracts, low stock warnings, and quick task triggers."),
        ("admin/manage-users.html", "User Account Provisioning", "Creates customer accounts and registers customer profiles resulting from face-to-face transactions."),
        ("admin/user-loan.html", "Direct Loan Assignment", "Form to assign a device contract to a customer, calculate amortization schedules, and link loan records directly to the customer account."),
        ("admin/inventory.html", "Inventory & AppleDB Engine", "Catalog control module integrated with AppleDB API for one-click specifications auto-fill, studio image mapping, and stock adjustments."),
        ("admin/loan-ledger.html", "Loan Ledger & Collections", "Central ledger to record amortization installment payments (Cash, GCash, Bank) and automatically recalculate outstanding balances."),
        ("admin/communication.html", "Live Support Inbox", "Central customer support desk allowing staff to review message threads and reply to inquiries in real time."),
        ("admin/reports.html", "Analytics & Reporting", "Generates visual revenue charts, loan performance statistics, and inventory velocity reports with CSV/PDF export capability."),
        ("admin/manage-admins.html", "Staff Administration", "Manages internal staff accounts, roles (Super Admin vs Staff), and revocations.")
    ]
    add_styled_table(doc, ["Page File", "Module Title", "Functional Description"], admin_pages, [1.6, 1.8, 3.4])

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ━━━ SECTION 5: DETAILED SUB-PROCESS FLOWS ━━━
    h5 = doc.add_heading(level=1)
    r_h5 = h5.add_run("5. Detailed Sub-Process Flowcharts")
    r_h5.font.color.rgb = RGBColor(30, 27, 75)

    # Sub-Flow 1
    h5_1 = doc.add_heading(level=2)
    r_h5_1 = h5_1.add_run("A. Customer Account Creation & Direct Loan Assignment (Face-to-Face Agreement)")
    r_h5_1.font.color.rgb = RGBColor(190, 18, 60)

    p_sf1 = doc.add_paragraph()
    p_sf1.add_run(
        "This sequence diagram models the exact transaction when a customer completes an agreement at the branch, "
        "and the Administrator registers their account and assigns their loan contract:"
    )

    seq_1 = """sequenceDiagram
    autonumber
    actor Customer as Customer (Walk-in)
    actor Admin as Admin Staff
    participant System as Admin Portal (manage-users.html)
    participant Database as PostgreSQL DB
    participant UserPortal as User Portal (account.html)

    Customer->>Admin: Completes physical agreement at branch / partner platform
    Customer->>Admin: Submits physical documents (Valid ID, Proof of Income)
    Admin->>System: Opens 'Create User & Assign Loan' Modal
    Admin->>System: Enters Full Name, Email, Contact, Address
    Admin->>System: Selects iPhone Model & Sets Terms (e.g., 24 Months, SRP, Downpayment)
    System->>System: Calculates Monthly Amortization & Schedule
    Admin->>System: Clicks 'Submit Account & Loan'
    System->>Database: INSERT into users table (Credentials & Profile)
    System->>Database: INSERT into loans table (Device, Principal, Monthly, Balance)
    System->>Database: UPDATE devices table (Deduct 1 from Stock)
    Database-->>System: Confirmation Success
    System-->>Admin: Displays Generated User Credentials
    Admin->>Customer: Hands over device & provides login credentials
    Customer->>UserPortal: Customer logs in to view active loan balance & schedule"""

    add_code_block(doc, seq_1, "Mermaid Sequence Diagram (Account Creation & Direct Loan Assignment)")

    # Sub-Flow 2
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h5_2 = doc.add_heading(level=2)
    r_h5_2 = h5_2.add_run("B. Payment Collection & Ledger Update Process")
    r_h5_2.font.color.rgb = RGBColor(190, 18, 60)

    p_sf2 = doc.add_paragraph()
    p_sf2.add_run(
        "This sequence diagram outlines the monthly amortization payment workflow, balance deduction, and immediate synchronization with the customer's portal:"
    )

    seq_2 = """sequenceDiagram
    autonumber
    actor Customer as Customer
    actor Admin as Admin Cashier
    participant Ledger as Loan Ledger (loan-ledger.html)
    participant Database as PostgreSQL DB
    participant CustomerView as Customer Portal (account.html)

    Customer->>Admin: Pays Monthly Installment (Cash, GCash, or Bank Transfer)
    Admin->>Ledger: Searches Customer Name or Loan Account #
    Ledger-->>Admin: Displays Outstanding Balance & Next Due Date
    Admin->>Ledger: Clicks 'Record Payment' & enters Amount Paid + Payment Mode
    Admin->>Ledger: Confirms Transaction
    Ledger->>Database: INSERT into payments table (Amount, Mode, Timestamp)
    Ledger->>Database: UPDATE loans table (Reduce remaining_balance, Advance due_date)
    Database-->>Ledger: Transaction Saved
    Ledger-->>Admin: Generates Official Payment Ledger Entry
    Admin->>Customer: Issues Official Acknowledgment Receipt
    Customer->>CustomerView: Refreshes account.html to see updated paid balance immediately"""

    add_code_block(doc, seq_2, "Mermaid Sequence Diagram (Payment Collection & Ledger Update)")

    # Sub-Flow 3
    doc.add_paragraph().paragraph_format.space_after = Pt(8)
    h5_3 = doc.add_heading(level=2)
    r_h5_3 = h5_3.add_run("C. AppleDB Inventory Integration & Auto-Population Process")
    r_h5_3.font.color.rgb = RGBColor(190, 18, 60)

    p_sf3 = doc.add_paragraph()
    p_sf3.add_run(
        "This flowchart models how administrators add or update devices using real-time AppleDB catalog search, "
        "eliminating manual entry errors and maintaining consistent official studio renders:"
    )

    mermaid_adb = """flowchart TD
    A([Admin in inventory.html]) --> B[Click 'Add Device']
    B --> C[Type Model Name in AppleDB Search e.g. iPhone 17]
    C --> D{Model Found in AppleDB?}
    D -- Yes --> E[Auto-Fetch: AppleDB Identifier, A-Series SoC, Colors, Official Studio Render]
    D -- No --> F[Manually Enter Specifications & Upload Custom Image]
    E --> G[Enter Unit SRP & Stock Count]
    F --> G
    G --> H[Select Availability Status: Available vs Upcoming]
    H --> I[Click 'Save Device']
    I --> J[Write Record to PostgreSQL Database]
    J --> K[Automatic Notification: Announcement Created if Upcoming or New Arrival]
    K --> L([Product Instantly Visible in User Catalog])"""

    add_code_block(doc, mermaid_adb, "Mermaid Diagram Code (AppleDB Inventory Integration)")

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # ━━━ SECTION 6: ROLES & PERMISSIONS MATRIX ━━━
    h6 = doc.add_heading(level=1)
    r_h6 = h6.add_run("6. Roles & Permissions Matrix")
    r_h6.font.color.rgb = RGBColor(30, 27, 75)

    matrix_data = [
        ("Browse Devices with AppleDB Studio Renders", "Read-Only (Available & Upcoming)", "Full Control (Add, Edit, Stock, Delete)"),
        ("View Device Hardware Specifications", "Read-Only (Display, SoC, Camera, Battery)", "Auto-Populate from AppleDB & Edit"),
        ("Customer Account Creation", "Login to assigned account", "Direct Creation & Profile Provisioning"),
        ("Device Loan Registration", "View Assigned Active Loan Only", "Direct Binding from Face-to-Face Agreement"),
        ("Payment Ledger & Balance Tracking", "View Personal Ledger & Payment History", "Record Payments, Manage Ledger Entries"),
        ("Customer Support Messaging", "Send Inquiries to Store Staff", "View Unified Inbox & Reply to All Clients"),
        ("Financial Reports & Sales Analytics", "No Access", "Full View, Filter, and Export (CSV/PDF)"),
        ("Admin Staff Accounts & Roles", "No Access", "Super Admin Full Control")
    ]
    add_styled_table(doc, ["Feature / System Module", "Customer (User Portal)", "Administrator (Admin Portal)"], matrix_data, [2.5, 2.1, 2.2])

    # Save to all target paths
    for p in output_paths:
        os.makedirs(os.path.dirname(p), exist_ok=True)
        doc.save(p)
        print(f"Successfully generated Word document at: {p}")

def add_styled_table(doc, headers, data_rows, col_widths):
    table = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].width = Inches(col_widths[i])
        set_cell_background(hdr_cells[i], "1E1B4B") # Deep Navy
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.bold = True
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, data in enumerate(data_rows):
        row_cells = table.rows[row_idx + 1].cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(data):
            row_cells[col_idx].width = Inches(col_widths[col_idx])
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=60, bottom=60, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            if col_idx == 0:
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)
            else:
                r.font.color.rgb = RGBColor(51, 65, 85)

def add_code_block(doc, code_text, caption):
    cap_p = doc.add_paragraph()
    cap_p.paragraph_format.space_before = Pt(8)
    cap_p.paragraph_format.space_after = Pt(3)
    r_cap = cap_p.add_run(f"► {caption}:")
    r_cap.font.bold = True
    r_cap.font.size = Pt(9.5)
    r_cap.font.color.rgb = RGBColor(30, 27, 75)

    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = tbl.rows[0].cells[0]
    c.width = Inches(6.8)
    set_cell_background(c, "F1F5F9") # Soft slate gray
    set_cell_margins(c, top=80, bottom=80, left=120, right=120)
    p = c.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = Pt(12)
    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(15, 23, 42)

if __name__ == "__main__":
    download_dir = r"C:\Users\Chase\Downloads"
    workspace_docs = r"c:\Users\Chase\Documents\my-supabase-project\docs"
    
    target_paths = [
        os.path.join(download_dir, "ADC_Gadgets_Flowcharts.docx"),
        os.path.join(workspace_docs, "ADC_Gadgets_Flowcharts.docx")
    ]
    create_flowchart_document(target_paths)
